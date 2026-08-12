import argparse
import csv
import glob
import json
import os
import re
import sys
from dataclasses import fields
from typing import List, Tuple, Dict, Any

import nibabel as nib
import numpy as np

sys.path.insert(0, os.getcwd())

from functions.nifti_area_sampler import AreaBandConfig, NiftiAreaSampler


# Label ids are scheme-specific: the same number is a different tissue in a
# different atlas, so there is deliberately no built-in default list. Every run
# states its labels in its own config, next to the legend they came from.
SEG_PATTERN_DEFAULT = "seg.nii.gz"


def _load_config(path: str) -> Tuple[AreaBandConfig, Dict[str, Any]]:
    with open(path, "r", encoding="utf-8-sig") as f:
        data = json.load(f)
    cfg_fields = {f.name for f in fields(AreaBandConfig)}
    cfg_kwargs = {k: v for k, v in data.items() if k in cfg_fields}
    if "file_path" not in cfg_kwargs:
        cfg_kwargs["file_path"] = ""
    if "out_dir" not in cfg_kwargs:
        cfg_kwargs["out_dir"] = ""
    extra = {k: v for k, v in data.items() if k not in cfg_fields}
    return AreaBandConfig(**cfg_kwargs), extra


def _cli_has(flag: str) -> bool:
    return flag in sys.argv


def _normalize_empty(value: str | None) -> str | None:
    if value is None:
        return None
    if not str(value).strip():
        return None
    if str(value).strip().lower() in {"none", "null"}:
        return None
    return value


def _validate_config(cfg: AreaBandConfig) -> None:
    if cfg.n < 1:
        raise ValueError("n must be >= 1")
    if not (0.0 < float(cfg.p) <= 1.0):
        raise ValueError("p must be in (0, 1]")
    NiftiAreaSampler._normalize_axis(cfg.axis)


def _legend_rows_csv(path: str):
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            yield row


def _legend_rows_xlsx(path: str):
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True).worksheets[0]
    for row in ws.iter_rows(values_only=True):
        yield ["" if c is None else str(c) for c in row]


def _legend_rows_itksnap(path: str):
    # ITK-SnAP label description: IDX  R  G  B  A  VIS  MSH  "LABEL"
    pattern = re.compile(r'^\s*(\d+)\s+(?:\S+\s+){6}"(.*)"\s*$')
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            m = pattern.match(line)
            if m:
                yield [m.group(1), m.group(2)]


def _load_label_legend(path: str | None) -> Dict[int, str]:
    """Map label id -> region name so errors and logs name tissues, not numbers.

    Handles the legend formats this project already carries: a two-column CSV
    (``info/dhcp-atlas-summary-info-19-labels.csv``), an ITK-SnAP ``.txt`` label
    description, and an ``.xlsx`` sheet (``assets/labels.xlsx``). An id cell may
    list several ids, e.g. ``"3, 4, 5,6"``, which all take the same name.
    Returns an empty map when no legend is configured.
    """
    if not path:
        return {}
    if not os.path.isfile(path):
        raise ValueError(f"label_legend not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        rows = _legend_rows_xlsx(path)
    elif ext == ".txt":
        rows = _legend_rows_itksnap(path)
    else:
        rows = _legend_rows_csv(path)

    legend: Dict[int, str] = {}
    for row in rows:
        cells = [str(c).strip() for c in row if str(c).strip()]
        if len(cells) < 2 or not re.fullmatch(r"[\d,\s]+", cells[0]):
            continue  # header or malformed row
        name = max(cells[1:], key=len)
        for i in re.findall(r"\d+", cells[0]):
            legend[int(i)] = name
    if not legend:
        raise ValueError(f"label_legend parsed no labels: {path}")
    return legend


def _describe_labels(ids, legend: Dict[int, str]) -> str:
    return ", ".join(f"{i} ({legend[i]})" if i in legend else str(i) for i in ids)


def _check_labels(file_path: str, labels: List[int] | None,
                  legend: Dict[int, str]) -> None:
    """Fail loudly on a label/volume mismatch, before any sampling happens.

    Two distinct mistakes are caught here. A requested id that the volume does
    not contain is always fatal. Separately, a volume carrying ids the legend
    does not describe means the legend does not belong to this dataset, which is
    the tell for a segmentation-scheme mix-up.

    Note the limit: when two schemes share an id range, a list written for one
    stays valid against the other and simply selects different tissue. Nothing
    here can detect that, which is why the resolved selection is always printed
    by name - reading it is the real check.
    """
    name = os.path.basename(file_path)
    data = np.asanyarray(nib.load(file_path).dataobj)
    present = {int(v) for v in np.rint(np.unique(data)) if int(v) != 0}

    missing = sorted(set(int(x) for x in (labels or [])) - present)
    if missing:
        raise ValueError(
            f"{name}: requested label(s) not in this volume: "
            f"{_describe_labels(missing, legend)}\n"
            f"  volume contains: {_describe_labels(sorted(present), legend)}\n"
            f"  Label ids are scheme-specific - check area_labels against this "
            f"dataset's legend."
        )

    if not legend:
        return

    undescribed = sorted(present - set(legend))
    if undescribed:
        raise ValueError(
            f"{name}: volume contains label(s) the legend does not describe: "
            f"{', '.join(str(i) for i in undescribed)}\n"
            f"  The legend does not match this volume - wrong label_legend for "
            f"this dataset?"
        )

    unseen = sorted(set(legend) - present - {0})
    if unseen:
        print(
            f"warning: {name}: legend describes label(s) absent from this volume: "
            f"{_describe_labels(unseen, legend)}",
            file=sys.stderr,
        )


def _find_seg(case_dir: str, pattern: str) -> str | None:
    """Locate the label volume inside a case folder.

    ``pattern`` is a glob, so datasets that do not use the original
    ``seg.nii.gz`` name work unchanged (the dHCP atlas ships, for example,
    ``tissue-t30.00_dhcp-19.nii.gz``).
    """
    matches = sorted(glob.glob(os.path.join(case_dir, pattern)))
    if len(matches) > 1:
        names = ", ".join(os.path.basename(m) for m in matches)
        raise ValueError(
            f"{case_dir}: seg pattern '{pattern}' matched {len(matches)} files "
            f"({names}); make the pattern more specific."
        )
    return matches[0] if matches else None


def _resolve_relative(path: str | None, config_path: str | None) -> str | None:
    """Resolve a path given in a config: as written first, then next to the config."""
    if not path or os.path.isabs(path) or os.path.isfile(path):
        return path
    if config_path:
        candidate = os.path.join(os.path.dirname(os.path.abspath(config_path)), path)
        if os.path.isfile(candidate):
            return candidate
    return path


def run_single(file_path: str, out_dir: str, axis: str, n: int, p: float,
               area_labels: List[int] | None, show_crosshair: bool, profile_plot: bool,
               save_png: bool = True, use_pial_overlay: bool = False,
               pial_lh_path: str | None = None, pial_rh_path: str | None = None,
               pial_space: str = "scanner", pial_line_thickness: float | None = None,
               label_legend: Dict[int, str] | None = None) -> dict:
    _check_labels(file_path, area_labels, label_legend or {})
    cfg = AreaBandConfig(
        file_path=file_path,
        out_dir=out_dir,
        axis=axis,
        n=n,
        p=p,
        save_png=save_png,
        profile_plot=profile_plot,
        show_crosshair=show_crosshair,
        area_labels=area_labels,
        use_pial_overlay=use_pial_overlay,
        pial_lh_path=pial_lh_path,
        pial_rh_path=pial_rh_path,
        pial_space=pial_space,
        pial_line_thickness=pial_line_thickness,
    )
    sampler = NiftiAreaSampler(cfg)
    return sampler.sample_band()


def run_batch(base_dir: str, out_base: str, axis: str, n: int, p: float,
              area_labels: List[int] | None, show_crosshair: bool, profile_plot: bool,
              save_png: bool = True, use_pial_overlay: bool = False,
              pial_lh_path: str | None = None, pial_rh_path: str | None = None,
              pial_space: str = "scanner", axis_subdir: bool = False,
              pial_line_thickness: float | None = None,
              seg_pattern: str = SEG_PATTERN_DEFAULT,
              label_legend: Dict[int, str] | None = None) -> list[dict]:
    results = []
    for name in sorted(os.listdir(base_dir)):
        d = os.path.join(base_dir, name)
        if not os.path.isdir(d):
            continue
        seg = _find_seg(d, seg_pattern)
        if seg is None:
            continue
        out_dir = os.path.join(out_base, name)
        if axis_subdir:
            out_dir = os.path.join(out_dir, f"axis_{axis}")
        os.makedirs(out_dir, exist_ok=True)
        res = run_single(
            seg,
            out_dir,
            axis,
            n,
            p,
            area_labels,
            show_crosshair,
            profile_plot,
            save_png,
            use_pial_overlay,
            pial_lh_path,
            pial_rh_path,
            pial_space,
            pial_line_thickness,
            label_legend=label_legend,
        )
        results.append({
            "case": name,
            "out_dir": out_dir,
            "x_max": res.get("x_max"),
            "f_max": res.get("f_max"),
            "left": res.get("left"),
            "right": res.get("right"),
            "n": len(res.get("positions") or []),
        })
        # Write a compact summary per case
        with open(os.path.join(out_dir, "area_band_run_summary.json"), "w", encoding="utf-8") as f:
            json.dump(results[-1], f, indent=2)
    if not results:
        raise ValueError(
            f"no case folder under {base_dir} contained a file matching "
            f"seg pattern '{seg_pattern}'"
        )
    return results


def main():
    ap = argparse.ArgumentParser(description="NIfTI area-band sampler runner")
    ap.add_argument("--config", help="Path to JSON config (AreaBandConfig fields + optional batch fields)")
    ap.add_argument("--all-axes", action="store_true", help="Run for all axes (x,y,z)")
    ap.add_argument("--file", help="Path to seg.nii.gz (single run)")
    ap.add_argument("--out", help="Output directory (single run)")
    ap.add_argument("--axis", default="z", choices=["x", "y", "z", "0", "1", "2"], help="Slicing axis")
    ap.add_argument("--n", type=int, default=10, help="Number of slices to sample")
    ap.add_argument("--p", type=float, default=0.8, help="Top-p threshold (0..1)")
    ap.add_argument("--labels", nargs="*", type=int, default=None, help="Label IDs to include in area/legend")
    ap.add_argument("--no-crosshair", action="store_true", help="Disable crosshair overlay")
    ap.add_argument("--no-profile-plot", action="store_true", help="Disable profile plotting")
    ap.add_argument("--no-png", action="store_true", help="Disable saving PNG images")

    ap.add_argument("--batch-dir", help="Base folder with multiple cases (each containing a label volume)")
    ap.add_argument("--batch-out", default="area_band_output", help="Output base for batch mode")
    ap.add_argument("--seg-glob", default=None,
                    help=f"Glob for the label volume inside each case folder "
                         f"(default '{SEG_PATTERN_DEFAULT}')")
    ap.add_argument("--label-legend", default=None,
                    help="Path to a label id -> region legend (.csv, .txt ITK-SnAP, or .xlsx). "
                         "Used to name labels in messages and to document the run.")
    ap.add_argument("--pial-lh", help="Left hemisphere pial surface path (Freesurfer .pial)")
    ap.add_argument("--pial-rh", help="Right hemisphere pial surface path (Freesurfer .pial)")
    ap.add_argument("--pial-space", default="scanner", choices=["scanner", "tkr"],
                    help="Pial coordinate space (scanner or tkr)")
    ap.add_argument("--pial-line-thickness", type=float, default=None,
                    help="Pial overlay line thickness in pixels (default 1)")
    ap.add_argument("--use-pial-overlay", action="store_true",
                    help="Enable pial overlay (auto-detect per case when --pial-lh/--pial-rh are not set)")
    ap.add_argument("--no-pial-overlay", action="store_true",
                    help="Disable pial overlay (overrides config)")
    ap.add_argument("--axis-subdir", action="store_true",
                    help="In batch mode, write results to per-axis subfolders (axis_x, axis_y, axis_z)")

    args = ap.parse_args()

    if args.config:
        cfg, extra = _load_config(args.config)

        axis = cfg.axis
        labels = list(cfg.area_labels) if cfg.area_labels is not None else None
        n = cfg.n
        p = cfg.p
        show_crosshair = cfg.show_crosshair
        profile_plot = cfg.profile_plot
        save_png = cfg.save_png
        use_pial_overlay = cfg.use_pial_overlay
        pial_lh_path = cfg.pial_lh_path
        pial_rh_path = cfg.pial_rh_path
        pial_space = cfg.pial_space
        pial_line_thickness = cfg.pial_line_thickness

        file_path = cfg.file_path
        out_dir = cfg.out_dir
        batch_dir = _normalize_empty(extra.get("batch_dir"))
        batch_out = _normalize_empty(extra.get("batch_out"))
        axis_subdir = bool(extra.get("axis_subdir", False))
        all_axes = bool(extra.get("all_axes", False))
        seg_pattern = _normalize_empty(extra.get("seg_pattern")) or SEG_PATTERN_DEFAULT
        legend_path = _normalize_empty(extra.get("label_legend"))

        if _cli_has("--axis"):
            axis = args.axis
        if _cli_has("--n"):
            n = args.n
        if _cli_has("--p"):
            p = args.p
        if _cli_has("--labels"):
            labels = args.labels
        if _cli_has("--seg-glob"):
            seg_pattern = args.seg_glob or SEG_PATTERN_DEFAULT
        if _cli_has("--label-legend"):
            legend_path = _normalize_empty(args.label_legend)
        if _cli_has("--file"):
            file_path = args.file
        if _cli_has("--out"):
            out_dir = args.out
        if _cli_has("--batch-dir"):
            batch_dir = _normalize_empty(args.batch_dir)
        if _cli_has("--batch-out"):
            batch_out = _normalize_empty(args.batch_out)
        if args.axis_subdir:
            axis_subdir = True
        if args.all_axes:
            all_axes = True
        if args.no_crosshair:
            show_crosshair = False
        if args.no_profile_plot:
            profile_plot = False
        if args.no_png:
            save_png = False
        if _cli_has("--pial-space"):
            pial_space = args.pial_space
        if _cli_has("--pial-line-thickness"):
            pial_line_thickness = args.pial_line_thickness
        if _cli_has("--pial-lh"):
            pial_lh_path = args.pial_lh
        if _cli_has("--pial-rh"):
            pial_rh_path = args.pial_rh
        if args.use_pial_overlay:
            use_pial_overlay = True
        if args.no_pial_overlay:
            use_pial_overlay = False

        if pial_lh_path or pial_rh_path:
            use_pial_overlay = True

        if axis in ("0", "1", "2"):
            axis = {"0": "x", "1": "y", "2": "z"}[axis]

        axes = ["x", "y", "z"] if all_axes else [axis]
        if all_axes and not axis_subdir:
            axis_subdir = True

        try:
            _validate_config(AreaBandConfig(file_path=file_path or "", out_dir=out_dir or "", axis=axis, n=n, p=p))
        except ValueError as exc:
            ap.error(str(exc))

        legend = _load_label_legend(_resolve_relative(legend_path, args.config))
        if labels:
            print(f"area_labels: {_describe_labels(sorted(labels), legend)}", file=sys.stderr)

        if batch_dir:
            if not os.path.isdir(batch_dir):
                ap.error(f"--batch-dir not found: {batch_dir}")
            if not batch_out:
                ap.error("Batch config missing batch_out (or provide --batch-out)")
            all_results = []
            for ax in axes:
                res = run_batch(
                    base_dir=batch_dir,
                    out_base=batch_out,
                    axis=ax,
                    n=n,
                    p=p,
                    area_labels=labels,
                    show_crosshair=show_crosshair,
                    profile_plot=profile_plot,
                    save_png=save_png,
                    use_pial_overlay=use_pial_overlay,
                    pial_lh_path=pial_lh_path,
                    pial_rh_path=pial_rh_path,
                    pial_space=pial_space,
                    axis_subdir=axis_subdir,
                    pial_line_thickness=pial_line_thickness,
                    seg_pattern=seg_pattern,
                    label_legend=legend,
                )
                all_results.append({"axis": ax, "results": res})
            print(json.dumps(all_results, indent=2))
            return 0

        if not file_path or not out_dir:
            ap.error("Config missing file_path/out_dir (or provide --file and --out)")

        single_results = {}
        for ax in axes:
            run_out = out_dir
            if axis_subdir:
                run_out = os.path.join(out_dir, f"axis_{ax}")
                os.makedirs(run_out, exist_ok=True)
            res = run_single(
                file_path=file_path,
                out_dir=run_out,
                axis=ax,
                n=n,
                p=p,
                area_labels=labels,
                show_crosshair=show_crosshair,
                profile_plot=profile_plot,
                save_png=save_png,
                use_pial_overlay=use_pial_overlay,
                pial_lh_path=pial_lh_path,
                pial_rh_path=pial_rh_path,
                pial_space=pial_space,
                pial_line_thickness=pial_line_thickness,
                label_legend=legend,
            )
            single_results[ax] = {k: res.get(k) for k in ("x_max", "f_max", "left", "right")}
        print(json.dumps(single_results, indent=2))
        return 0

    axis = args.axis
    if axis in ("0", "1", "2"):
        axis = {"0": "x", "1": "y", "2": "z"}[axis]

    labels = args.labels
    legend = _load_label_legend(_resolve_relative(_normalize_empty(args.label_legend), None))
    seg_pattern = args.seg_glob or SEG_PATTERN_DEFAULT
    if labels:
        print(f"area_labels: {_describe_labels(sorted(labels), legend)}", file=sys.stderr)

    if args.batch_dir:
        if not os.path.isdir(args.batch_dir):
            ap.error(f"--batch-dir not found: {args.batch_dir}")
        axes = ["x", "y", "z"] if args.all_axes else [axis]
        try:
            _validate_config(AreaBandConfig(file_path=args.file or "", out_dir=args.out or "", axis=axis, n=args.n, p=args.p))
        except ValueError as exc:
            ap.error(str(exc))
        all_results = []
        for ax in axes:
            res = run_batch(
                base_dir=args.batch_dir,
                out_base=args.batch_out,
                axis=ax,
                n=args.n,
                p=args.p,
                area_labels=labels,
                show_crosshair=not args.no_crosshair,
                profile_plot=not args.no_profile_plot,
                save_png=not args.no_png,
                use_pial_overlay=bool(args.pial_lh or args.pial_rh or args.use_pial_overlay),
                pial_lh_path=args.pial_lh,
                pial_rh_path=args.pial_rh,
                pial_space=args.pial_space,
                axis_subdir=(args.axis_subdir or args.all_axes),
                pial_line_thickness=args.pial_line_thickness,
                seg_pattern=seg_pattern,
                label_legend=legend,
            )
            all_results.append({"axis": ax, "results": res})
        print(json.dumps(all_results, indent=2))
        return 0

    if not args.file or not args.out:
        ap.error("For single run, provide --file and --out (or use --batch-dir)")

    axes = ["x", "y", "z"] if args.all_axes else [axis]
    try:
        _validate_config(AreaBandConfig(file_path=args.file, out_dir=args.out, axis=axis, n=args.n, p=args.p))
    except ValueError as exc:
        ap.error(str(exc))
    if len(axes) > 1:
        results = {}
        for ax in axes:
            out_dir = os.path.join(args.out, f"axis_{ax}")
            os.makedirs(out_dir, exist_ok=True)
            res = run_single(
                file_path=args.file,
                out_dir=out_dir,
                axis=ax,
                n=args.n,
                p=args.p,
                area_labels=labels,
                show_crosshair=not args.no_crosshair,
                profile_plot=not args.no_profile_plot,
                save_png=not args.no_png,
                use_pial_overlay=bool(args.pial_lh or args.pial_rh or args.use_pial_overlay),
                pial_lh_path=args.pial_lh,
                pial_rh_path=args.pial_rh,
                pial_space=args.pial_space,
                pial_line_thickness=args.pial_line_thickness,
                label_legend=legend,
            )
            results[ax] = {k: res.get(k) for k in ("x_max", "f_max", "left", "right")}
        print(json.dumps(results, indent=2))
        return 0

    res = run_single(
        file_path=args.file,
        out_dir=args.out,
        axis=axis,
        n=args.n,
        p=args.p,
        area_labels=labels,
        show_crosshair=not args.no_crosshair,
        profile_plot=not args.no_profile_plot,
        save_png=not args.no_png,
        use_pial_overlay=bool(args.pial_lh or args.pial_rh or args.use_pial_overlay),
        pial_lh_path=args.pial_lh,
        pial_rh_path=args.pial_rh,
        pial_space=args.pial_space,
        pial_line_thickness=args.pial_line_thickness,
        label_legend=legend,
    )
    print(json.dumps({k: res.get(k) for k in ("x_max", "f_max", "left", "right")}, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ValueError as exc:
        print(f"area_band_cli.py: error: {exc}", file=sys.stderr)
        raise SystemExit(2)
